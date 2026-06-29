"""Tests for the multi-factor equity screen.

Written test-first against the precise scoring contract in the task spec.
"""
import math
import os
import sys

import numpy as np
import pandas as pd
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import factor_screen as fs  # noqa: E402


# --------------------------------------------------------------------------- #
# 1. Scoring math on a tiny hand-built fixture
# --------------------------------------------------------------------------- #
def _frame(rows):
    return pd.DataFrame(rows)


def test_higher_earnings_yield_gives_higher_value_score():
    df = _frame([
        {"ticker": f"T{i}", "name": f"n{i}", "region": "US", "earnings_yield": ey}
        for i, ey in enumerate([0.01, 0.03, 0.05, 0.07, 0.09])
    ])
    out = fs.score_universe(df, scale=10, rank_scope="global")
    scores = out.sort_values("earnings_yield" if "earnings_yield" in out else "ticker")
    vals = out.set_index("ticker")["Value"]
    # Ascending earnings_yield -> non-decreasing Value score, strictly up overall.
    seq = [vals[f"T{i}"] for i in range(5)]
    assert seq == sorted(seq)
    assert seq[0] < seq[-1]


def test_higher_vol_gives_lower_lowvol_score():
    df = _frame([
        {"ticker": f"T{i}", "name": f"n{i}", "region": "US", "vol_annual": v}
        for i, v in enumerate([0.10, 0.20, 0.30, 0.40, 0.50])
    ])
    out = fs.score_universe(df, scale=10, rank_scope="global").set_index("ticker")
    seq = [out.loc[f"T{i}", "Low Vol"] for i in range(5)]
    # Direction flip: higher vol -> lower score (strictly decreasing here).
    assert seq == sorted(seq, reverse=True)
    assert seq[0] > seq[-1]


def test_value_composite_averages_available_components_ignoring_nan():
    # Two of three Value components present for the target row; the third is NaN
    # and must be skipped (not imputed). With a controlled cross-section we can
    # predict the percentile of each present component and hence the composite.
    df = _frame([
        {"ticker": "A", "name": "a", "region": "US",
         "earnings_yield": 0.10, "book_to_price": 0.10, "ebit_ev_yield": np.nan},
        {"ticker": "B", "name": "b", "region": "US",
         "earnings_yield": 0.05, "book_to_price": 0.20, "ebit_ev_yield": 0.10},
        {"ticker": "C", "name": "c", "region": "US",
         "earnings_yield": 0.01, "book_to_price": 0.30, "ebit_ev_yield": 0.20},
    ])
    out = fs.score_universe(df, scale=10, rank_scope="global").set_index("ticker")
    # A: earnings_yield is top (pct 1.0), book_to_price is bottom (pct 1/3).
    # Composite percentile = (1.0 + 1/3) / 2 = 0.6667 -> ceil(6.667)=7.
    assert out.loc["A", "Value"] == 7
    # Coverage for A counts Value only (no other factor columns present).
    assert out.loc["A", "Coverage"] == 1


def test_coverage_counts_when_whole_factor_missing():
    # Provide Value + Quality columns but NO Momentum/LowVol/Skin columns.
    df = _frame([
        {"ticker": "A", "name": "a", "region": "US", "earnings_yield": 0.08, "roe": 0.2},
        {"ticker": "B", "name": "b", "region": "US", "earnings_yield": 0.04, "roe": 0.1},
    ])
    out = fs.score_universe(df, scale=10, rank_scope="global").set_index("ticker")
    assert out.loc["A", "Coverage"] == 2  # Value + Quality only
    assert pd.isna(out.loc["A", "Momentum"])
    assert pd.isna(out.loc["A", "Low Vol"])
    assert pd.isna(out.loc["A", "Skin in Game"])


def test_ceil_mapping_endpoints():
    assert fs.score_from_percentile(1.0, 10) == 10   # top percentile -> K
    assert fs.score_from_percentile(0.0001, 10) == 1  # bottom -> 1
    assert fs.score_from_percentile(0.5, 10) == 5
    assert fs.score_from_percentile(0.51, 10) == 6
    assert fs.score_from_percentile(float("nan"), 10) is None
    # With n >= K, the literal bottom/top of a cross-section hit 1 and K.
    df = _frame([
        {"ticker": f"T{i}", "name": f"n{i}", "region": "US", "earnings_yield": float(i)}
        for i in range(10)
    ])
    out = fs.score_universe(df, scale=10, rank_scope="global").set_index("ticker")
    assert out.loc["T9", "Value"] == 10
    assert out.loc["T0", "Value"] == 1


def test_derivation_from_vendor_primitives():
    df = _frame([
        {"ticker": "A", "name": "a", "region": "US", "pe": 10.0},   # ey = 0.10
        {"ticker": "B", "name": "b", "region": "US", "pe": -5.0},   # negative -> NaN
        {"ticker": "C", "name": "c", "region": "US", "pe": 20.0},   # ey = 0.05
    ])
    derived = fs.derive_metrics(df)
    assert math.isclose(derived.loc[0, "earnings_yield"], 0.10)
    assert pd.isna(derived.loc[1, "earnings_yield"])
    assert math.isclose(derived.loc[2, "earnings_yield"], 0.05)


# --------------------------------------------------------------------------- #
# 2. Region-neutral ranking
# --------------------------------------------------------------------------- #
def test_region_scope_keeps_each_region_near_midpoint():
    counts = {"US": 400, "DE": 400, "CN": 400, "IN": 400}
    df = fs.make_demo(counts, seed=7)
    out = fs.score_universe(df, scale=10, rank_scope="region")
    midpoint = (10 + 1) / 2  # 5.5
    for region, g in out.groupby("region"):
        for factor in fs.FACTOR_NAMES:
            m = g[factor].mean()
            assert abs(m - midpoint) < 1.0, (region, factor, m)


# --------------------------------------------------------------------------- #
# 3. Demo at scale -> workbook with expected sheets and zero formula errors
# --------------------------------------------------------------------------- #
def test_demo_at_scale_workbook(tmp_path):
    from openpyxl import load_workbook

    counts = {"US": 5000, "DE": 500, "CN": 300, "IN": 500}
    n = sum(counts.values())
    df = fs.make_demo(counts, seed=1)
    assert len(df) == n
    scored = fs.score_universe(df, scale=10, rank_scope="region")
    out = tmp_path / "scale.xlsx"
    fs.write_workbook(scored, str(out), scale=10)

    wb = load_workbook(str(out))
    assert set(["Scores", "Top 50", "Region summary", "Methodology"]).issubset(set(wb.sheetnames))
    ws = wb["Scores"]
    # Weights row (1) + header row (2) + n data rows.
    assert ws.max_row == n + 2
    assert wb["Top 50"].max_row == 50 + 1

    # Prefer a real LibreOffice recalc; fall back to formula-string validation
    # when LO cannot load files in this environment (the spec's documented
    # fallback). Either way, assert zero formula errors.
    errors = None
    if fs.libreoffice_available():
        try:
            recalced = fs.recalc_with_libreoffice(str(out))
            errors = fs.find_error_cells(recalced)
        except Exception:
            errors = None
    if errors is None:
        errors = fs.validate_formulas(str(out), scale=10)
    assert errors == [], errors[:10]


# --------------------------------------------------------------------------- #
# 4. CSV ingestion with only a subset of metric columns
# --------------------------------------------------------------------------- #
def test_csv_subset_scores_only_present_factors(tmp_path):
    csv = tmp_path / "u.csv"
    pd.DataFrame([
        {"ticker": "A", "name": "Aco", "region": "US", "pe": 10.0, "roe": 0.20},
        {"ticker": "B", "name": "Bco", "region": "US", "pe": 25.0, "roe": 0.05},
        {"ticker": "C", "name": "Cco", "region": "US", "pe": 15.0, "roe": 0.12},
    ]).to_csv(csv, index=False)

    df = fs.load_csv(str(csv))
    out = fs.score_universe(df, scale=10, rank_scope="global").set_index("ticker")
    # Value (from pe) and Quality (from roe) scored; the rest blank.
    assert out.loc["A", "Value"] == 10  # cheapest pe
    assert out.loc["A", "Quality"] == 10  # highest roe
    for blank in ["Momentum", "Low Vol", "Skin in Game"]:
        assert out[blank].isna().all()
    assert (out["Coverage"] == 2).all()


def test_load_excel_input(tmp_path):
    # The loader accepts an .xlsx input too (format chosen by extension).
    xlsx = tmp_path / "u.xlsx"
    pd.DataFrame([
        {"ticker": "A", "name": "Aco", "region": "US", "pe": 10.0, "roe": 0.20},
        {"ticker": "B", "name": "Bco", "region": "US", "pe": 25.0, "roe": 0.05},
    ]).to_excel(xlsx, index=False)
    df = fs.load_csv(str(xlsx))
    assert list(df["ticker"]) == ["A", "B"]
    out = fs.score_universe(df, scale=10, rank_scope="global").set_index("ticker")
    assert out.loc["A", "Value"] == 10


def test_committed_sample_workbook_scores():
    # The example input shipped in examples/ must load and score cleanly.
    sample = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                          "examples", "sample_universe.xlsx")
    if not os.path.exists(sample):
        pytest.skip("sample workbook not present")
    df = fs.load_csv(sample)
    out = fs.score_universe(df, scale=10, rank_scope="region")
    assert len(out) == 13
    assert out["Coverage"].max() == 5


# --------------------------------------------------------------------------- #
# 5. Weights change Weighted but not Total
# --------------------------------------------------------------------------- #
def test_weights_change_weighted_not_total():
    df = fs.make_demo({"US": 200, "DE": 100}, seed=3)
    base = fs.score_universe(df, scale=10, rank_scope="region",
                             weights=[1, 1, 1, 1, 1]).set_index("ticker")
    bumped = fs.score_universe(df, scale=10, rank_scope="region",
                               weights=[2, 1, 1, 1, 1]).set_index("ticker")
    # Total identical between the two runs.
    pd.testing.assert_series_equal(base["Total"], bumped["Total"], check_names=False)
    # With all-1 weights, Weighted == Total.
    eq = base.dropna(subset=["Weighted", "Total"])
    assert np.allclose(eq["Weighted"], eq["Total"])
    # Bumping the Value weight changes Weighted for rows that have a Value score.
    has_value = bumped["Value"].notna()
    assert not np.allclose(bumped.loc[has_value, "Weighted"],
                           base.loc[has_value, "Weighted"])
